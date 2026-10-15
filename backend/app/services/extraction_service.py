import os
import csv
from io import StringIO
from typing import List, Dict, Any
from pydantic import BaseModel
from app.core.logging import logger

class ExtractedSection(BaseModel):
    text: str
    location: str
    metadata: Dict[str, Any] = {}

class ExtractionResult(BaseModel):
    sections: List[ExtractedSection]
    total_characters: int
    format_metadata: Dict[str, Any] = {}

class ExtractionService:
    @staticmethod
    def extract_txt(file_path: str) -> ExtractionResult:
        encodings = ["utf-8", "utf-16", "latin-1"]
        content = ""
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue

        sections = []
        # Split by double newlines or sections to retain meaningful paragraph locations
        raw_paras = [p.strip() for p in content.split("\n\n") if p.strip()]
        if not raw_paras and content.strip():
            raw_paras = [content.strip()]

        for idx, para in enumerate(raw_paras, 1):
            sections.append(ExtractedSection(
                text=para,
                location=f"Section {idx}",
                metadata={"section_index": idx}
            ))

        return ExtractionResult(
            sections=sections,
            total_characters=len(content),
            format_metadata={"line_count": len(content.splitlines()), "paragraphs": len(raw_paras)}
        )

    @staticmethod
    def extract_pdf(file_path: str) -> ExtractionResult:
        sections = []
        total_chars = 0
        total_pages = 0

        # Try pypdf
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            total_pages = len(reader.pages)
            for page_idx, page in enumerate(reader.pages, 1):
                page_text = page.extract_text() or ""
                page_text = page_text.strip()
                if page_text:
                    total_chars += len(page_text)
                    sections.append(ExtractedSection(
                        text=page_text,
                        location=f"Page {page_idx}",
                        metadata={"page": page_idx, "total_pages": total_pages}
                    ))
        except Exception as e:
            logger.error(f"pypdf extraction failed for {file_path}: {e}")
            raise ValueError(f"Failed to extract PDF: {str(e)}")

        return ExtractionResult(
            sections=sections,
            total_characters=total_chars,
            format_metadata={"pages": total_pages}
        )

    @staticmethod
    def extract_docx(file_path: str) -> ExtractionResult:
        try:
            import docx
            doc = docx.Document(file_path)
            sections = []
            total_chars = 0
            
            # Paragraphs
            for p_idx, p in enumerate(doc.paragraphs, 1):
                text = p.text.strip()
                if text:
                    total_chars += len(text)
                    sections.append(ExtractedSection(
                        text=text,
                        location=f"Paragraph {p_idx}",
                        metadata={"paragraph_index": p_idx}
                    ))
            
            # Tables
            for t_idx, table in enumerate(doc.tables, 1):
                table_rows = []
                for row_idx, row in enumerate(table.rows, 1):
                    row_cells = [cell.text.strip() for cell in row.cells]
                    row_str = " | ".join(row_cells)
                    if row_str.strip():
                        table_rows.append(row_str)
                if table_rows:
                    table_text = "\n".join(table_rows)
                    total_chars += len(table_text)
                    sections.append(ExtractedSection(
                        text=table_text,
                        location=f"Table {t_idx}",
                        metadata={"table_index": t_idx, "rows_count": len(table_rows)}
                    ))

            return ExtractionResult(
                sections=sections,
                total_characters=total_chars,
                format_metadata={"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
            )
        except Exception as e:
            logger.error(f"DOCX extraction failed for {file_path}: {e}")
            raise ValueError(f"Failed to extract DOCX: {str(e)}")

    @staticmethod
    def extract_csv(file_path: str) -> ExtractionResult:
        sections = []
        total_chars = 0
        try:
            encodings = ["utf-8", "latin-1"]
            for enc in encodings:
                try:
                    with open(file_path, "r", encoding=enc) as f:
                        reader = csv.reader(f)
                        rows = list(reader)
                    break
                except UnicodeDecodeError:
                    continue

            if not rows:
                return ExtractionResult(sections=[], total_characters=0, format_metadata={})

            header = rows[0]
            header_str = ", ".join(header)

            for row_idx, row in enumerate(rows[1:], 2):
                if not any(row):
                    continue
                # Format row with column names for full row context
                row_parts = []
                for col_idx, val in enumerate(row):
                    col_name = header[col_idx] if col_idx < len(header) else f"Column_{col_idx+1}"
                    row_parts.append(f"{col_name}: {val.strip()}")
                row_text = "; ".join(row_parts)
                total_chars += len(row_text)
                sections.append(ExtractedSection(
                    text=row_text,
                    location=f"Row {row_idx}",
                    metadata={"row": row_idx, "columns": header}
                ))

            return ExtractionResult(
                sections=sections,
                total_characters=total_chars,
                format_metadata={"total_rows": len(rows), "columns": header}
            )
        except Exception as e:
            logger.error(f"CSV extraction failed for {file_path}: {e}")
            raise ValueError(f"Failed to extract CSV: {str(e)}")

    @staticmethod
    def extract_xlsx(file_path: str) -> ExtractionResult:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sections = []
            total_chars = 0

            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue

                # First non-empty row as header
                header = None
                data_start_idx = 0
                for idx, r in enumerate(rows):
                    if any(r):
                        header = [str(c).strip() if c is not None else f"Col_{i+1}" for i, c in enumerate(r)]
                        data_start_idx = idx + 1
                        break

                if not header:
                    continue

                for row_offset, row in enumerate(rows[data_start_idx:], data_start_idx + 1):
                    if not any(row):
                        continue
                    row_parts = []
                    for col_idx, val in enumerate(row):
                        if val is not None and str(val).strip():
                            col_name = header[col_idx] if col_idx < len(header) else f"Col_{col_idx+1}"
                            row_parts.append(f"{col_name}: {str(val).strip()}")
                    if row_parts:
                        row_text = "; ".join(row_parts)
                        total_chars += len(row_text)
                        sections.append(ExtractedSection(
                            text=row_text,
                            location=f"Sheet: {sheet_name}, Row {row_offset}",
                            metadata={"sheet": sheet_name, "row": row_offset}
                        ))

            return ExtractionResult(
                sections=sections,
                total_characters=total_chars,
                format_metadata={"sheets": wb.sheetnames}
            )
        except Exception as e:
            logger.error(f"XLSX extraction failed for {file_path}: {e}")
            raise ValueError(f"Failed to extract XLSX: {str(e)}")

    @classmethod
    def extract(cls, file_path: str, file_type: str) -> ExtractionResult:
        ext = file_type.lower()
        if not ext.startswith("."):
            ext = f".{ext}"

        if ext == ".pdf":
            return cls.extract_pdf(file_path)
        elif ext == ".docx":
            return cls.extract_docx(file_path)
        elif ext == ".txt":
            return cls.extract_txt(file_path)
        elif ext == ".csv":
            return cls.extract_csv(file_path)
        elif ext == ".xlsx":
            return cls.extract_xlsx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")
